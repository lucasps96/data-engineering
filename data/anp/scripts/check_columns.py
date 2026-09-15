"""
Valida se todos os arquivos .xlsx baixados têm o mesmo cabeçalho de colunas
na aba MUNICIPIOS.

Lida com arquivos que o openpyxl não consegue abrir diretamente (formato
OOXML "strict" com referências externas quebradas) fazendo fallback
automático para conversão via LibreOffice headless.

Uso (de dentro da pasta data/anp/, via toolbox):
    docker-run.sh scripts/check_columns.py

Requisitos:
    pip install openpyxl --break-system-packages
    (libreoffice precisa estar disponível no container/host para o fallback)
"""

import os
import subprocess
import tempfile
import warnings
from collections import defaultdict

import openpyxl

warnings.filterwarnings("ignore")

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DADOS_DIR = os.path.join(SCRIPT_DIR, "..", "dados")

SHEET_NAME = "MUNICIPIOS"
HEADER_ROW_INDEX = 9  # linha 10 (0-indexed) — confirmado na inspeção manual


def try_open_workbook(path: str):
    """
    Tenta abrir o workbook diretamente. Se falhar (arquivo com formato
    'strict' OOXML malformado), converte via LibreOffice headless para
    um .xlsx padrão e tenta de novo.
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if wb.sheetnames:
            return wb, False  # abriu direto, sem precisar de fallback
    except Exception:
        pass

    # Fallback: converte via LibreOffice headless
    with tempfile.TemporaryDirectory() as tmpdir:
        result = subprocess.run(
            [
                "libreoffice", "--headless", "--calc", "--convert-to", "xlsx",
                "--outdir", tmpdir, path,
            ],
            capture_output=True,
            text=True,
            timeout=90,
        )
        converted_path = os.path.join(tmpdir, os.path.basename(path))
        if not os.path.exists(converted_path):
            raise RuntimeError(
                f"LibreOffice não conseguiu converter o arquivo. "
                f"stderr: {result.stderr[:300]}"
            )
        wb = openpyxl.load_workbook(converted_path, read_only=True, data_only=True)
        return wb, True  # precisou de fallback


def get_header(path: str):
    """Retorna a tupla de cabeçalho da aba MUNICIPIOS, e se precisou de fallback."""
    wb, used_fallback = try_open_workbook(path)
    try:
        if SHEET_NAME not in wb.sheetnames:
            return None, used_fallback, f"aba '{SHEET_NAME}' não encontrada (abas: {wb.sheetnames})"
        ws = wb[SHEET_NAME]
        rows = list(ws.iter_rows(min_row=HEADER_ROW_INDEX + 1, max_row=HEADER_ROW_INDEX + 1, values_only=True))
        if not rows:
            return None, used_fallback, "linha de cabeçalho vazia"
        return rows[0], used_fallback, None
    finally:
        wb.close()


def main():
    files = sorted(f for f in os.listdir(DADOS_DIR) if f.lower().endswith(".xlsx"))
    print(f"Verificando {len(files)} arquivos em {DADOS_DIR}\n")

    headers_seen = defaultdict(list)  # header -> [filenames]
    fallback_used = []
    errors = []

    for i, filename in enumerate(files, 1):
        path = os.path.join(DADOS_DIR, filename)
        print(f"[{i}/{len(files)}] {filename}", end=" ... ")
        try:
            header, used_fallback, err = get_header(path)
            if err:
                print(f"ERRO: {err}")
                errors.append((filename, err))
                continue
            headers_seen[header].append(filename)
            if used_fallback:
                fallback_used.append(filename)
                print("OK (via LibreOffice fallback)")
            else:
                print("OK")
        except Exception as e:
            print(f"FALHA: {e}")
            errors.append((filename, str(e)))

    print("\n" + "=" * 70)
    print("RESUMO")
    print("=" * 70)

    print(f"\nTotal de arquivos: {len(files)}")
    print(f"Lidos com sucesso: {sum(len(v) for v in headers_seen.values())}")
    print(f"Precisaram de fallback (LibreOffice): {len(fallback_used)}")
    print(f"Falharam completamente: {len(errors)}")

    print(f"\nVariações de cabeçalho encontradas: {len(headers_seen)}")
    for idx, (header, filenames) in enumerate(headers_seen.items(), 1):
        print(f"\n  Variação {idx} ({len(filenames)} arquivos):")
        print(f"    {header}")
        if len(filenames) <= 5:
            print(f"    Arquivos: {filenames}")
        else:
            print(f"    Exemplos: {filenames[:3]} ... (+{len(filenames)-3} outros)")

    if fallback_used:
        print(f"\nArquivos que precisaram de conversão via LibreOffice ({len(fallback_used)}):")
        for f in fallback_used:
            print(f"  - {f}")

    if errors:
        print(f"\nArquivos com erro ({len(errors)}):")
        for f, e in errors:
            print(f"  - {f}: {e}")


if __name__ == "__main__":
    main()
