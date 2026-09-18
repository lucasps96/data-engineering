"""
Extração de uma aba específica de um arquivo "resumo semanal" da ANP
(bronze) para um Parquet limpo (silver intermediário).

Generalizado para funcionar com qualquer uma das 5 abas do arquivo
(MUNICIPIOS, ESTADOS, REGIOES, CAPITAIS, BRASIL) — cada aba tem seu
próprio conjunto de colunas geográficas (ex.: REGIOES usa "REGIÃO" em
vez de "ESTADO"/"MUNICÍPIO"), então o cabeçalho é lido dinamicamente
a partir do próprio arquivo, em vez de assumir uma lista fixa de colunas.

Este script cuida apenas da extração/reparo do Excel — não faz
transformação de negócio, que fica a cargo da etapa seguinte, em Spark.

Uso (de dentro de data/anp/, via toolbox):
    docker-run.sh scripts/anp_extract.py --arquivo dados/arquivo.xlsx --aba MUNICIPIOS

    # Saída padrão: parquet/<aba>/<mesmo nome>.parquet
    # Para especificar destino:
    docker-run.sh scripts/anp_extract.py --arquivo dados/arquivo.xlsx --aba ESTADOS --saida parquet/saida.parquet

Requisitos:
    pip install openpyxl pandas pyarrow --break-system-packages
    (libreoffice precisa estar disponível no container, para o fallback)
"""

import argparse
import os
import subprocess
import sys
import tempfile
import warnings

import openpyxl
import pandas as pd

warnings.filterwarnings("ignore")

# Todas as 5 abas seguem o mesmo padrão de layout: título e metadados
# nas primeiras linhas, cabeçalho na linha 10 (índice 9). Isso foi
# confirmado manualmente para MUNICIPIOS; assume-se o mesmo padrão para
# as demais, dado que são geradas pelo mesmo processo de exportação —
# a validação de cabeçalho abaixo funciona como uma checagem dessa
# suposição (se a linha 10 não tiver texto nenhum, o erro aparece aqui).
HEADER_ROW_INDEX = 9

SHEETS_DISPONIVEIS = ["MUNICIPIOS", "ESTADOS", "REGIOES", "CAPITAIS", "BRASIL"]


def try_open_workbook(path: str):
    """
    Tenta abrir o workbook diretamente. Se falhar (arquivo com formato
    'strict' OOXML malformado), converte via LibreOffice headless e tenta
    de novo.

    Retorna (workbook, usou_fallback: bool).
    """
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if wb.sheetnames:
            return wb, False
    except Exception:
        pass

    with tempfile.TemporaryDirectory() as tmpdir:
        result = subprocess.run(
            [
                "libreoffice", "--headless", "--calc", "--convert-to", "xlsx",
                "--outdir", tmpdir, path,
            ],
            capture_output=True,
            text=True,
            timeout=90,
        )
        converted_path = os.path.join(tmpdir, os.path.basename(path))
        if not os.path.exists(converted_path):
            raise RuntimeError(
                f"LibreOffice não conseguiu converter o arquivo. "
                f"stderr: {result.stderr[:300]}"
            )
        fd, persistent_path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        with open(converted_path, "rb") as src, open(persistent_path, "wb") as dst:
            dst.write(src.read())

    try:
        wb = openpyxl.load_workbook(persistent_path, read_only=True, data_only=True)
    finally:
        os.remove(persistent_path)
    return wb, True


def extract_sheet(path: str, sheet_name: str) -> tuple[pd.DataFrame, bool]:
    """
    Extrai uma aba específica de um arquivo bronze e retorna um DataFrame
    com o cabeçalho lido dinamicamente da linha 10 (as colunas variam
    conforme a aba — ex.: REGIOES não tem MUNICÍPIO).
    """
    if sheet_name not in SHEETS_DISPONIVEIS:
        raise ValueError(f"Aba '{sheet_name}' desconhecida. Opções: {SHEETS_DISPONIVEIS}")

    wb, used_fallback = try_open_workbook(path)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"Aba '{sheet_name}' não encontrada em {os.path.basename(path)}. "
                f"Abas disponíveis: {wb.sheetnames}"
            )
        ws = wb[sheet_name]

        header_row = next(
            ws.iter_rows(min_row=HEADER_ROW_INDEX + 1, max_row=HEADER_ROW_INDEX + 1, values_only=True)
        )
        # Mantém só as colunas com nome real, descartando células vazias
        # no final da linha (mesmo problema visto em alguns arquivos de MUNICIPIOS)
        header = [str(h).strip() for h in header_row if h is not None]
        n_cols = len(header)

        if n_cols == 0:
            raise ValueError(
                f"Linha de cabeçalho vazia na aba '{sheet_name}' de {os.path.basename(path)} "
                f"— verifique se HEADER_ROW_INDEX ainda é válido para esta aba."
            )

        data_rows = []
        for row in ws.iter_rows(min_row=HEADER_ROW_INDEX + 2, values_only=True):
            if all(c is None for c in row):
                continue
            data_rows.append(row[:n_cols])

        df = pd.DataFrame(data_rows, columns=header)
        return df, used_fallback
    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--arquivo", required=True, help="Caminho do .xlsx de entrada (bronze)")
    parser.add_argument(
        "--aba", default="MUNICIPIOS", choices=SHEETS_DISPONIVEIS,
        help="Aba a extrair (default: MUNICIPIOS)"
    )
    parser.add_argument("--saida", default=None, help="Caminho do .parquet de saída (default: parquet/<aba>/<nome>.parquet)")
    args = parser.parse_args()

    if not os.path.exists(args.arquivo):
        print(f"ERRO: arquivo não encontrado: {args.arquivo}", file=sys.stderr)
        sys.exit(1)

    if args.saida is None:
        base_name = os.path.splitext(os.path.basename(args.arquivo))[0]
        script_dir = os.path.dirname(os.path.abspath(__file__))
        out_dir = os.path.join(script_dir, "..", "parquet", args.aba.lower())
        os.makedirs(out_dir, exist_ok=True)
        saida = os.path.join(out_dir, base_name + ".parquet")
    else:
        saida = args.saida
        os.makedirs(os.path.dirname(os.path.abspath(saida)) or ".", exist_ok=True)

    print(f"Lendo: {args.arquivo} (aba: {args.aba})")
    df, used_fallback = extract_sheet(args.arquivo, args.aba)

    if used_fallback:
        print("  (precisou de fallback via LibreOffice)")

    df["ARQUIVO_ORIGEM"] = os.path.basename(args.arquivo)

    # coerce_timestamps="us": o pandas/pyarrow grava timestamp em nanossegundos
    # por padrão, formato que o Spark (Parquet reader) não sabe ler
    # (erro "Illegal Parquet type: INT64 (TIMESTAMP(NANOS,false))").
    # Forçar microssegundos garante compatibilidade sem perda de precisão
    # relevante para este dado (granularidade é de dias, não nanossegundos).
    df.to_parquet(saida, index=False, coerce_timestamps="us", allow_truncated_timestamps=True)

    print(f"Linhas extraídas: {len(df)}")
    print(f"Colunas: {list(df.columns)}")
    print(f"Salvo em: {saida}")

    print("\nAmostra (5 primeiras linhas):")
    print(df.head().to_string())


if __name__ == "__main__":
    main()
